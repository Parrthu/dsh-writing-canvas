#!/usr/bin/env python3
"""写作画布的 DOCX 格式引擎：生成 + 回读校验。

用法：
    python3 docx_tool.py <job.json>

job.json：
    {
      "mode": "generate" | "verify" | "inspect",
      "content": "Markdown 正文",
      "spec": { ...格式规格... },
      "specId": "gongwen-gb9704",
      "title": "文档标题（可选，缺省取正文第一个标题）",
      "outPath": "/输出/路径.docx",
      "markdown": true
    }

输出：stdout 上一个 JSON 对象，形如
    { "ok": true, "file": "...", "bytes": 1234,
      "verification": { "passed": true, "checks": [ {name, expected, actual, ok} ] },
      "warnings": [ "..." ] }

设计原则：**绝不谎报成功**。生成之后一定重新打开文件逐项核对，
任何一项对不上就把 ok 置为 false 并在 checks 里写清楚差在哪。
"""

from __future__ import annotations

import json
import os
import re
import sys
from typing import Any

try:
    import docx
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor
except ImportError as exc:  # pragma: no cover - 环境缺库时给出可读原因
    print(
        json.dumps(
            {
                "ok": False,
                "error": "python-docx-missing",
                "message": f"缺少 python-docx：{exc}",
            },
            ensure_ascii=False,
        )
    )
    sys.exit(0)


# --------------------------------------------------------------------------
# Markdown 子集解析
# --------------------------------------------------------------------------

# 支持的 Markdown 子集：标题、引用、无序/有序列表、分隔线、空行分段，
# 以及行内 **加粗**、*斜体*、`代码`。其余语法按普通文本处理（不猜测、不报错）。

INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")


def parse_inline(text: str) -> list[tuple[str, dict]]:
    """把一行文本切成 (片段, 样式) 序列。"""
    parts: list[tuple[str, dict]] = []
    for chunk in INLINE_PATTERN.split(text):
        if chunk == "":
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            parts.append((chunk[2:-2], {"bold": True}))
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            parts.append((chunk[1:-1], {"italic": True}))
        elif chunk.startswith("`") and chunk.endswith("`") and len(chunk) > 2:
            parts.append((chunk[1:-1], {"code": True}))
        else:
            parts.append((chunk, {}))
    return parts or [("", {})]


def parse_blocks(content: str) -> list[dict]:
    """把 Markdown 内容解析成块序列。

    @param content: Markdown 正文。
    @returns: [{ kind, text, level }]，kind ∈ title/h1/h2/h3/quote/ul/ol/body/hr
    """
    blocks: list[dict] = []
    for raw in content.replace("\r\n", "\n").split("\n"):
        line = raw.rstrip()
        stripped = line.strip()

        if stripped == "":
            blocks.append({"kind": "blank", "text": ""})
            continue
        if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", stripped):
            blocks.append({"kind": "hr", "text": ""})
            continue

        heading = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            blocks.append({"kind": f"h{level}", "text": heading.group(2).strip()})
            continue

        if stripped.startswith("> "):
            blocks.append({"kind": "quote", "text": stripped[2:].strip()})
            continue

        bullet = re.match(r"^[-*+]\s+(.*)$", stripped)
        if bullet:
            blocks.append({"kind": "ul", "text": bullet.group(1).strip()})
            continue

        numbered = re.match(r"^\d+[.)]\s+(.*)$", stripped)
        if numbered:
            blocks.append({"kind": "ol", "text": numbered.group(1).strip()})
            continue

        blocks.append({"kind": "body", "text": stripped})
    return blocks


# --------------------------------------------------------------------------
# 字体与段落工具
# --------------------------------------------------------------------------


def apply_run_font(run, spec: dict, *, bold: bool | None = None) -> None:
    """把中西文字体、字号、加粗写到一个 run 上。"""
    font = run.font
    font.name = spec.get("fontAscii") or "Times New Roman"
    size_pt = spec.get("sizePt")
    if size_pt:
        font.size = Pt(float(size_pt))
    if bold is not None:
        font.bold = bool(bold)
    # 关键：中文字体必须写 w:eastAsia，否则 Word 会退回默认宋体。
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    east = spec.get("fontEastAsia")
    if east:
        rfonts.set(qn("w:eastAsia"), east)
    rfonts.set(qn("w:ascii"), spec.get("fontAscii") or "Times New Roman")
    rfonts.set(qn("w:hAnsi"), spec.get("fontAscii") or "Times New Roman")


def apply_paragraph_format(paragraph, spec: dict) -> None:
    """把行距、对齐、首行缩进、段间距写到一个段落上。"""
    fmt = paragraph.paragraph_format

    rule = spec.get("lineRule")
    spacing_pt = spec.get("lineSpacingPt")
    multiple = spec.get("lineSpacingMultiple")
    if rule == "exact" and spacing_pt:
        fmt.line_spacing = Pt(float(spacing_pt))
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    elif multiple:
        fmt.line_spacing = float(multiple)
        fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE

    align = spec.get("align")
    if align == "center":
        fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "right":
        fmt.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    if spec.get("spaceBeforePt"):
        fmt.space_before = Pt(float(spec["spaceBeforePt"]))
    if spec.get("spaceAfterPt"):
        fmt.space_after = Pt(float(spec["spaceAfterPt"]))

    # 首行缩进按「字符」写，这样换字号也不会走样（Word 的 firstLineChars）。
    chars = spec.get("firstLineIndentChars")
    if chars:
        ppr = paragraph._p.get_or_add_pPr()
        ind = ppr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            ppr.append(ind)
        ind.set(qn("w:firstLineChars"), str(int(float(chars) * 100)))
        size_pt = spec.get("sizePt") or 12
        ind.set(qn("w:firstLine"), str(int(float(size_pt) * 20 * float(chars))))


def add_page_number_footer(section, spec: dict) -> None:
    """在页脚居中放页码；公文规格会在数字两侧加一字线。"""
    footer = section.footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = paragraph.paragraph_format
    fmt.line_spacing = Pt(float(spec.get("sizePt") or 14) * 1.2)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY

    if spec.get("dash"):
        apply_run_font(paragraph.add_run("— "), spec)
    run = paragraph.add_run()
    apply_run_font(run, spec)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._element.append(fld_begin)
    run._element.append(instr)
    run._element.append(fld_end)
    if spec.get("dash"):
        apply_run_font(paragraph.add_run(" —"), spec)


# --------------------------------------------------------------------------
# 生成
# --------------------------------------------------------------------------


def build_document(content: str, spec: dict, title: str | None) -> Document:
    """按规格把 Markdown 内容生成一个 Word 文档。"""
    document = Document()

    # 默认样式也设成正文规格，避免未显式设置的文字走回 Word 默认字体。
    normal = document.styles["Normal"]
    normal.font.name = spec["body"].get("fontAscii") or "Times New Roman"
    normal.font.size = Pt(float(spec["body"].get("sizePt") or 12))
    normal_rpr = normal.element.get_or_add_rPr()
    normal_fonts = normal_rpr.find(qn("w:rFonts"))
    if normal_fonts is None:
        normal_fonts = OxmlElement("w:rFonts")
        normal_rpr.append(normal_fonts)
    normal_fonts.set(qn("w:eastAsia"), spec["body"].get("fontEastAsia") or "宋体")

    section = document.sections[0]
    page = spec.get("page", {})
    if page.get("widthMm"):
        section.page_width = Mm(float(page["widthMm"]))
    if page.get("heightMm"):
        section.page_height = Mm(float(page["heightMm"]))
    for key, attr in (
        ("marginTopMm", "top_margin"),
        ("marginBottomMm", "bottom_margin"),
        ("marginLeftMm", "left_margin"),
        ("marginRightMm", "right_margin"),
    ):
        if page.get(key) is not None:
            setattr(section, attr, Mm(float(page[key])))

    blocks = parse_blocks(content)

    # 标题处理：显式传入的标题优先，并且要把正文里**重复的那个一级标题去掉**，
    # 否则标题会出现两次，后续按位置抽样校验正文时也会取错段落。
    doc_title = title
    if doc_title is None:
        for index, block in enumerate(blocks):
            if block["kind"] == "h1":
                doc_title = block["text"]
                blocks = blocks[:index] + blocks[index + 1 :]
                break
    else:
        for index, block in enumerate(blocks):
            if block["kind"] == "blank":
                continue
            if block["kind"] == "h1" and block["text"].strip() == doc_title.strip():
                blocks = blocks[:index] + blocks[index + 1 :]
            break

    if doc_title:
        paragraph = document.add_paragraph()
        apply_paragraph_format(paragraph, {**spec.get("title", {}), "firstLineIndentChars": 0})
        for text, style in parse_inline(doc_title):
            run = paragraph.add_run(text)
            apply_run_font(run, spec.get("title", {}), bold=style.get("bold"))

    style_map = {
        "h1": spec.get("heading1", spec["body"]),
        "h2": spec.get("heading2", spec["body"]),
        "h3": spec.get("heading3", spec["body"]),
        "quote": spec.get("quote", spec["body"]),
    }

    for block in blocks:
        kind = block["kind"]
        if kind == "blank":
            continue
        if kind == "hr":
            paragraph = document.add_paragraph()
            apply_paragraph_format(paragraph, {**spec["body"], "firstLineIndentChars": 0})
            apply_run_font(paragraph.add_run("—" * 18), spec["body"])
            continue

        paragraph = document.add_paragraph()
        if kind in ("ul", "ol"):
            paragraph.style = document.styles["List Bullet" if kind == "ul" else "List Number"]
            apply_paragraph_format(paragraph, {**spec["body"], "firstLineIndentChars": 0})
            run_spec = spec["body"]
        else:
            run_spec = style_map.get(kind, spec["body"])
            apply_paragraph_format(paragraph, run_spec)

        for text, style in parse_inline(block["text"]):
            run = paragraph.add_run(text)
            bold = style.get("bold")
            if style.get("bold") is None and run_spec.get("bold") is not None:
                bold = run_spec.get("bold")
            apply_run_font(run, run_spec, bold=bold)
            if style.get("italic"):
                run.font.italic = True
            if style.get("code"):
                run.font.name = "Consolas"
                rpr = run._element.get_or_add_rPr()
                rfonts = rpr.find(qn("w:rFonts"))
                if rfonts is not None:
                    rfonts.set(qn("w:ascii"), "Consolas")
                    rfonts.set(qn("w:hAnsi"), "Consolas")

    if spec.get("pageNumber"):
        add_page_number_footer(section, spec["pageNumber"])

    return document


# --------------------------------------------------------------------------
# 回读校验
# --------------------------------------------------------------------------


def east_asia_font_of(run) -> str | None:
    """读出一个 run 实际写入的中文字体名。"""
    rpr = run._element.find(qn("w:rPr"))
    if rpr is None:
        return None
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        return None
    return rfonts.get(qn("w:eastAsia"))


def _close(a: Any, b: Any, tolerance: float = 0.6) -> bool:
    try:
        return abs(float(a) - float(b)) <= tolerance
    except (TypeError, ValueError):
        return False


def verify_document(path: str, spec: dict) -> dict:
    """重新打开生成的文件，逐项核对规格是否真的写进去了。"""
    checks: list[dict] = []

    def check(name: str, expected: Any, actual: Any, ok: bool) -> None:
        checks.append({"name": name, "expected": expected, "actual": actual, "ok": bool(ok)})

    document = Document(path)
    section = document.sections[0]
    page = spec.get("page", {})

    def mm(value) -> float | None:
        return None if value is None else round(value.mm, 1)

    for key, attr, label in (
        ("marginTopMm", "top_margin", "页边距·上"),
        ("marginBottomMm", "bottom_margin", "页边距·下"),
        ("marginLeftMm", "left_margin", "页边距·左"),
        ("marginRightMm", "right_margin", "页边距·右"),
    ):
        if page.get(key) is None:
            continue
        actual = mm(getattr(section, attr))
        check(f"{label}(mm)", page[key], actual, _close(actual, page[key]))

    paragraphs = [p for p in document.paragraphs if p.text.strip() != ""]
    body = spec["body"]

    # 标题段落：居中 + 字号 + 中文字体
    if paragraphs and spec.get("title"):
        title_paragraph = paragraphs[0]
        title_spec = spec["title"]
        alignment = title_paragraph.paragraph_format.alignment
        check(
            "标题·居中对齐",
            "CENTER",
            str(alignment),
            alignment is not None and "CENTER" in str(alignment),
        )
        runs = [r for r in title_paragraph.runs if r.text.strip() != ""]
        if runs:
            size = runs[0].font.size
            actual_pt = None if size is None else round(size.pt, 2)
            check("标题·字号(pt)", title_spec.get("sizePt"), actual_pt, _close(actual_pt, title_spec.get("sizePt")))
            check("标题·中文字体", title_spec.get("fontEastAsia"), east_asia_font_of(runs[0]), east_asia_font_of(runs[0]) == title_spec.get("fontEastAsia"))

    # 正文段落：**按字体定位真正的正文段**，而不是盲取第二段。
    # 标题、各级标题、引用各有自己的字体，取错了会得到一堆假失败。
    body_paragraphs = [p for p in paragraphs[1:] if p.text.strip() != ""]
    body_font = body.get("fontEastAsia")
    body_samples: list = []
    font_histogram: dict[str, int] = {}
    for paragraph in body_paragraphs:
        runs = [r for r in paragraph.runs if r.text.strip() != ""]
        if not runs:
            continue
        font = east_asia_font_of(runs[0])
        if font:
            font_histogram[font] = font_histogram.get(font, 0) + 1
        if font == body_font:
            body_samples.append(paragraph)

    check(
        "正文段落数(使用正文字体)",
        ">=1",
        len(body_samples),
        len(body_samples) >= 1,
    )

    sample = body_samples[0] if body_samples else None
    if sample is not None:
        runs = [r for r in sample.runs if r.text.strip() != ""]
        size = runs[0].font.size
        actual_pt = None if size is None else round(size.pt, 2)
        check("正文·字号(pt)", body.get("sizePt"), actual_pt, _close(actual_pt, body.get("sizePt")))
        check(
            "正文·中文字体",
            body_font,
            east_asia_font_of(runs[0]),
            east_asia_font_of(runs[0]) == body_font,
        )

        fmt = sample.paragraph_format
        spacing = fmt.line_spacing
        if body.get("lineRule") == "exact" and body.get("lineSpacingPt"):
            actual_pt = None if spacing is None else round(spacing.pt, 2)
            check("正文·行距(pt)", body.get("lineSpacingPt"), actual_pt, _close(actual_pt, body.get("lineSpacingPt")))
            rule = str(fmt.line_spacing_rule)
            check("正文·行距规则", "EXACTLY", rule, "EXACTLY" in rule)
        elif body.get("lineSpacingMultiple"):
            actual = None if spacing is None else round(float(spacing), 3)
            check("正文·行距倍数", body.get("lineSpacingMultiple"), actual, _close(actual, body.get("lineSpacingMultiple"), 0.05))

    # 首行缩进（按字符）
    if sample is not None and body.get("firstLineIndentChars"):
        ppr = sample._p.find(qn("w:pPr"))
        ind = None if ppr is None else ppr.find(qn("w:ind"))
        chars = None if ind is None else ind.get(qn("w:firstLineChars"))
        expected = str(int(float(body["firstLineIndentChars"]) * 100))
        check("正文·首行缩进(字符×100)", expected, chars, chars == expected)

    failed = [item for item in checks if not item["ok"]]
    return {
        "passed": len(failed) == 0,
        "checks": checks,
        "failed": len(failed),
        "total": len(checks),
        "fontHistogram": font_histogram,
    }


# --------------------------------------------------------------------------
# 入口
# --------------------------------------------------------------------------


def main() -> None:
    if len(sys.argv) < 2:
        print(json.dumps({"ok": False, "error": "missing-job-file"}, ensure_ascii=False))
        return

    with open(sys.argv[1], "r", encoding="utf-8") as handle:
        job = json.load(handle)

    mode = job.get("mode", "generate")
    spec = job.get("spec") or {}
    out_path = job.get("outPath") or ""

    if mode == "inspect":
        report = verify_document(job["content"], spec)
        print(json.dumps({"ok": True, "verification": report}, ensure_ascii=False))
        return

    if mode == "verify":
        report = verify_document(out_path, spec)
        print(json.dumps({"ok": report["passed"], "verification": report}, ensure_ascii=False))
        return

    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    document = build_document(job.get("content", ""), spec, job.get("title"))
    document.save(out_path)

    report = verify_document(out_path, spec)
    warnings: list[str] = []
    if not report["passed"]:
        warnings.append("生成后的回读校验未全部通过，详见 checks。")
    warnings.append(
        "校验核对的是写入文件里的字体名称、字号与行距；实际渲染效果取决于查看文档的机器是否安装了这些字体。"
    )

    print(
        json.dumps(
            {
                "ok": report["passed"],
                "file": out_path,
                "bytes": os.path.getsize(out_path),
                "specId": job.get("specId"),
                "verification": report,
                "warnings": warnings,
                "python-docx": getattr(docx, "__version__", "unknown"),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
